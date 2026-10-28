"""
core/excel_manager.py - Excel Başvuran Yönetimi
plan.md Bölüm 4'teki Excel şemasını okur, doğrular ve şablon üretir.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
from typing import List, Dict, Any, Optional

from config import DATA_DIR

COLUMNS = [
    # Girdi Sütunları
    "id", "ad", "soyad", "pasaport_no", "dogum_tarihi", "uyruk", "telefon",
    "vfs_email", "vfs_sifre", "imap_app_sifre", "imap_sunucu", "merkez_tercih",
    "kategori", "tarih_baslangic", "tarih_bitis",
    # Bot Çıktı Sütunları
    "durum", "randevu_tarihi", "randevu_saati", "merkez", "referans_no",
    "deneme_sayisi", "son_hata", "guncelleme_zamani"
]

SAMPLE_ROW = [
    1, "Ahmet", "Yılmaz", "U12345678", "1992-04-15", "TR", "+905321112233",
    "ornek_vfs_hesap@gmail.com", "GucluSifre.2026", "abcd efgh ijkl mnop", "imap.gmail.com",
    "IST, BUR, ANK", "Turistik", "2026-11-01", "2026-12-31",
    "PENDING", "", "", "", "", 0, "", ""
]


class ExcelManager:
    """Excel başvuran dosyalarını işleme ve şablon oluşturma."""

    @staticmethod
    def get_default_excel_path() -> Path:
        return DATA_DIR / "applicants.xlsx"

    @classmethod
    def create_template(cls, output_path: Optional[Path] = None) -> Path:
        """Kullanıcının doldurması için görsel olarak biçimlendirilmiş bir Excel şablonu oluşturur."""
        target_path = output_path or (DATA_DIR / "vfs_bulgaristan_sablon.xlsx")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Başvuranlar"

        # Başlık Stili
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        output_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")

        ws.append(COLUMNS)

        for col_idx, col_name in enumerate(COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.alignment = center_align
            # İlk 15 sütun kullanıcı girdisi (mavi), sonrası bot çıktısı (yeşil/teal)
            cell.fill = header_fill if col_idx <= 15 else output_fill

        # Örnek satır ekle
        ws.append(SAMPLE_ROW)

        # Sütun genişliklerini otomatik ayarla
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        target_path = Path(target_path)
        wb.save(target_path)
        return target_path

    @classmethod
    def parse_excel(cls, file_path: Path | str) -> List[Dict[str, Any]]:
        """Yüklenen Excel dosyasını okur ve JSON sözlük listesine dönüştürür."""
        file_path = Path(file_path)
        if not file_path.exists():
            return []

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        header_row = [str(c).strip().lower() if c else "" for c in rows[0]]
        applicants = []

        for row_idx, row in enumerate(rows[1:], 2):
            if not any(row):  # Tamamen boş satır
                continue

            item = {}
            for col_idx, col_name in enumerate(COLUMNS):
                # Başlık eşleşmesi varsa
                if col_name in header_row:
                    idx = header_row.index(col_name)
                    val = row[idx] if idx < len(row) else ""
                else:
                    val = row[col_idx] if col_idx < len(row) else ""
                
                item[col_name] = str(val).strip() if val is not None else ""

            # En azından ad veya email varsa listeye ekle
            if item.get("ad") or item.get("vfs_email"):
                applicants.append(item)

        return applicants

    @classmethod
    def save_applicants(cls, file_path: Path, applicants: List[Dict[str, Any]]) -> None:
        """Başvuranlar listesini Excel dosyasına kaydeder."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Başvuranlar"

        ws.append(COLUMNS)
        for app in applicants:
            row = [app.get(col, "") for col in COLUMNS]
            ws.append(row)

        wb.save(file_path)
