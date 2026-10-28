"""
bot/excel_io.py - Excel = görünüm katmanı. Atomik yazma: tmp → os.replace.
"""
from __future__ import annotations
import asyncio
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .config import CFG
from .db import DB, INPUT_COLS, OUTPUT_COLS

HEADERS = INPUT_COLS + OUTPUT_COLS


def read_inputs() -> list[dict]:
    path = Path(CFG["excel"]["path"])
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True)
    ws = wb[CFG["excel"]["sheet"]] if CFG["excel"]["sheet"] in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        row = {header[i]: raw[i] for i in range(min(len(header), len(raw))) if header[i] in INPUT_COLS}
        if row.get("id") and row.get("vfs_email"):
            out.append(row)
    return out


def write_outputs() -> None:
    """DB → Excel. Geçici dosyaya yaz, sonra atomik olarak yerine koy."""
    path = Path(CFG["excel"]["path"])
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = CFG["excel"]["sheet"]
    ws.append(HEADERS)
    for a in DB.applicants():
        ws.append([a[c] if c in a.keys() else "" for c in HEADERS])

    wb.save(tmp)
    os.replace(tmp, path)  # atomik


def create_template() -> None:
    path = Path(CFG["excel"]["path"])
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = CFG["excel"]["sheet"]
    ws.append(HEADERS)
    ws.append([
        1, "AHMET", "YILMAZ", "U12345678", "1990-05-12", "TR", "+905550000000",
        "ornek@gmail.com", "GucluSifre.2026", "gmail_uygulama_sifresi", "imap.gmail.com",
        "IST,ANK,EDR", "SCHENGEN", "2026-11-01", "2026-12-31", "", "PENDING",
    ])
    wb.save(path)


def sync_inputs_to_db() -> int:
    n = 0
    for row in read_inputs():
        DB.upsert_applicant(row)
        n += 1
    return n


async def excel_sync_loop() -> None:
    while True:
        try:
            await asyncio.to_thread(sync_inputs_to_db)
            await asyncio.to_thread(write_outputs)
        except Exception as exc:
            DB.log("ERROR", f"excel sync: {exc}")
        await asyncio.sleep(float(CFG["excel"]["sync_interval_s"]))
